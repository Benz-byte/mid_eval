from pathlib import Path
import ast
import hashlib
import re
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DOCX = OUT / 'Auto Scheduler Code Guide.docx'

DESCRIPTIONS = {
    'App.tsx': 'Owns the active CSV, fingerprint, scoped admin events, tab selection, archive preview, and cross-screen callbacks. Its upload, event commit, and archive-close handlers connect UI actions to persistence.',
    'main.tsx': 'Applies the saved theme and mounts App in React StrictMode. StrictMode can rerun development effects to expose unsafe side effects; it is not a second solver or database.',
    'main.ts': 'Starts the packaged Python backend, opens the Electron window, chooses development versus packaged content, and coordinates normal shutdown.',
    'backendProcess.ts': 'Stops only the spawned backend. On Windows it waits for taskkill to terminate that PID and its descendants; failures are reported rather than silently clearing ownership.',
    'environment.ts': 'Reads NODE_ENV to decide whether Electron should use the development URL or packaged resources.',
    'preload.ts': 'Exposes a small electron bridge containing the local Flask URL. The interface does not receive unrestricted Node access.',
    'ScheduleCalendar.tsx': 'Coordinates daily and weekly room views, conflict warnings, uploads, date navigation, filters, duty assignment, event-assistant selection, and relievers. Availability is calculated before saving and relevant actions recalculate it at confirmation.',
    'ScheduleCalendar.css': 'Styles the room timetable, absolute-positioned event blocks, conflict colors, responsive controls, dialogs, and focus highlights. Styling changes do not change the solver model.',
    'ScheduleDateNavigation.tsx': 'Combines previous and next navigation with the date heading and date-picker interaction. It passes a chosen date to the parent rather than importing a new schedule.',
    'ScheduleDateNavigation.css': 'Keeps the date navigation control consistently sized and styles its picker and navigation buttons.',
    'ScheduleFilter.tsx': 'Maintains room and instructor selections in a filter dialog and passes the chosen sets back to the calendar.',
    'ScheduleFilter.css': 'Styles filter choices, dialog layout, and selected states without changing stored schedule records.',
    'CalendarZoomControls.tsx': 'Provides shared minus, reset, and plus controls with 75–200 percent bounds and 25-point steps.',
    'CalendarZoomControls.css': 'Styles the small zoom control group used in both timetable screens.',
    'StudentAssistantPanel.tsx': 'Coordinates assistant profiles, parsing personal schedules, editing and deleting students, settings, full or incremental solving, selection, week navigation, and the workload popup.',
    'StudentAssistant.css': 'Styles the assistant sidebar, profile forms, settings, weekly calendar blocks, duty states, and animated Creating schedule indicator. The indicator is not a measured solver percentage.',
    'AssistantWeeklyCalendar.tsx': 'Draws one student’s personal classes, regular duties, assigned events, relieved duties, and replacement duties for a chosen week. It also handles duty details, View and Remove, and read-only preview behavior.',
    'WeeklyWorkloadSummary.tsx': 'Renders the roster-wide ranked workload report and its ascending or descending sort control. The calculations are supplied by weeklyWorkload.ts.',
    'AdminEventsPanel.tsx': 'Keeps per-date room and time selections, previews exam dates, validates the booking form, manages existing bookings, and requests confirmations for occupied selections or editing.',
    'AutoExpandingTextarea.tsx': 'Resizes title and description textareas to their content so long text does not remain trapped in a single-line control.',
    'AdminEventsPanel.css': 'Styles event forms, room availability states, selected dates and times, warnings, and management controls.',
    'bookingForm.ts': 'Converts stored booking cards into editable form defaults, infers date repetition, and supplies date-label helpers. The saved cards remain the authoritative schedule data.',
    'ArchiveManager.tsx': 'Lists local snapshots first, refreshes shared records separately, saves locally before closing, and offers view, delete, or conflict-only rename actions.',
    'ArchivePreview.tsx': 'Displays an explicit archive snapshot in read-only mode and switches between its schedule and assistant previews without activating it as current work.',
    'ArchivedAssistantPanel.tsx': 'Allows choosing an archived assistant, moving through weeks, zooming, and reading the archived workload ranking. It does not run the solver or save edits.',
    'Archives.css': 'Styles archive dialogs, action rows, archive-list cards, read-only banners, and responsive layouts.',
    'ThemeToggle.tsx': 'Displays the accessible light and dark toggle and responds to theme notifications.',
    'theme.ts': 'Applies the chosen theme to the document and notifies mounted controls. Storage access is delegated to themeStorage.ts.',
    'theme.css': 'Defines dark and light CSS variables used across feature styles, including stronger light-mode grid and button colors.',
    'App.css': 'Styles the application shell, tabs, shared navigation, and top-level layout.',
    'index.css': 'Defines global base typography, buttons, form controls, and other shared browser presentation defaults.',
    'types.ts': 'Defines CalendarEvent, UploadedAssistant, booking form data, metadata, and other shared TypeScript contracts. Types describe shapes; they do not replace runtime validation of external data.',
    'apiClient.ts': 'Builds local backend requests, parses JSON, preserves response status in BackendError, and provides timed polling subscriptions.',
    'scheduleApi.ts': 'Transfers the active imported schedule through the shared schedule endpoint. isCloudConfigured is a compatibility flag here; the backend determines whether credentials are configured.',
    'scheduleParser.ts': 'Reads CSV quoting and XLS/XLSX cell values in the interface, then sends normalized row arrays to the local Python parser.',
    'studentAssistantApi.ts': 'Defines settings, assignments, relievers, solver results, and shared payload contracts. It sends solver requests and verifies backend acknowledgements for scoped and versioned saving.',
    'adminEventApi.ts': 'Reads event snapshots with deleted IDs, saves event cards, distinguishes permanent HTTP 410 deletion from temporary failure, and checks protected-delete acknowledgements.',
    'archiveApi.ts': 'Transfers archive records and translates duplicate-name responses into ArchiveNameConflictError so storage can preserve the local snapshot and continue other sync work.',
    'scheduleStorage.ts': 'Loads and saves the active imported schedule with compatibility validation for rooms, times, metadata, and TBA fields.',
    'studentAssistantStorage.ts': 'Stores and activates per-CSV rosters and result caches, builds manual-assignment results, advances actual edit revisions, and merges remote records without replacing newer local data.',
    'adminEventStorage.ts': 'Stores event cards and permanently known deleted IDs, filters stale records, and emits a local notification when deletions are learned.',
    'localFirstSync.ts': 'Persists pending active-schedule, event, and assistant changes; serializes each flush; and acknowledges only the versions actually sent. It is a retry mechanism, not a guarantee of transactional multi-device editing.',
    'archiveStorage.ts': 'Uses IndexedDB transactions for immutable snapshots, local duplicate-name checks, deletion states, cloud retries, and conflict-only renaming.',
    'archiveTypes.ts': 'Defines archive data contracts and creates a cloned snapshot scoped to the active CSV, including viewing context.',
    'archiveWorkflow.ts': 'Waits for a local archive commit before invoking the optional close action. It contains no prerequisite cloud request.',
    'preferenceStorage.ts': 'Persists the selected tab and calendar date and restores valid defaults when saved preference values are absent or invalid.',
    'themeStorage.ts': 'Stores the device’s light or dark preference and falls back to dark if storage is unavailable.',
    'assistantSchedule.ts': 'Plans safe assistant edits and removes one assistant’s assignments without clearing the full result. It updates related labels, checks changed personal schedules, and recalculates totals.',
    'bookingEdit.ts': 'Matches existing booking cards to edited cards without duplicating assignments and revalidates assistants for changed slots.',
    'bookingSeries.ts': 'Derives a booking-series key by removing the trailing numeric indices and date suffix.',
    'examPeriod.ts': 'Recognizes exam marker and booking IDs, encodes schedule ownership, checks dates, and creates internal exam marker records.',
    'eventSchedule.ts': 'Filters events by schedule fingerprint and known room names while preserving exam markers and safely hiding unscoped legacy events.',
    'weeklyWorkload.ts': 'Calculates effective duty, event, and reliever minutes for each student in a selected week, subtracts merged coverage, and sorts the ranking.',
    'dateFormatter.ts': 'Converts local date input values, finds Monday week starts, formats week labels, and matches recurring weekday codes.',
    'timeFormatter.ts': 'Converts HH:MM input to minutes and formats numeric minutes and workload durations for people.',
    'assistantNameFormatter.ts': 'Formats available first, middle, and last name fields for full-name display while retaining fallback support for older stored labels.',
    'archiveName.ts': 'Normalizes archive names for case-insensitive, repeated-space-insensitive comparisons.',
    'roomScheduleExcel.ts': 'Creates per-room Monday–Friday XLSX worksheets with exact time boundaries, merged blocks, XML styles, safe names, and red conflict groups, then triggers a browser download.',
    'app.py': 'Creates Flask, enables the prototype’s CORS configuration, registers all route groups, and converts repository RuntimeErrors to HTTP 503. The main entry point listens on local port 5000.',
    'schedule_parser.py': 'Contains legacy, official-template, personal-schedule, and course-offering parsers. It normalizes headers, times, weekdays, sections, credits, metadata, TBA entries, room-only rows, and time-only rows.',
    'student_assistant_solver.py': 'Builds OR-Tools CP-SAT decision variables, hard constraints, and weighted preferences; supports incremental filling; and returns assignments, totals, diagnostics, and applied settings.',
    'client.py': 'Finds the development or frozen environment configuration, checks whether Supabase is configured, performs REST requests with a timeout, and surfaces service errors.',
    'schedule_repository.py': 'Stores and loads the shared active schedule at ccs-main and supports older list-based versus newer envelope-based JSON storage.',
    'student_assistant_repository.py': 'Merges per-CSV rosters, unlinked legacy assistants, results, and revisions into the shared assistant cache envelope before saving.',
    'admin_event_repository.py': 'Paginates event and deletion reads, translates application fields to database columns, blocks deleted IDs, and uses the protected delete RPC.',
    'archive_repository.py': 'Loads complete snapshots and tombstones, creates immutable archives idempotently, detects normalized duplicate names, and marks deletions separately from active schedules.',
    'schedule_service.py': 'Validates the shared schedule request and delegates loading or saving to the schedule repository.',
    'student_assistant_service.py': 'Connects assistant storage requests to validation and repository calls and solver requests to OR-Tools.',
    'admin_event_service.py': 'Validates event updates, delegates protected deletion, and builds sync snapshots filtered against deleted IDs.',
    'archive_service.py': 'Validates snapshot payloads and UUIDs before invoking archive creation, listing, or deletion.',
    'conflict_detector.py': 'Provides a small room-and-interval overlap predicate. It does not itself compare recurring weekdays or dates and is not the live calendar’s main conflict implementation.',
    'vacant_room_finder.py': 'Uses the simple backend overlap helper to return rooms without a conflicting supplied record. No current Flask route calls this helper; live room-picking logic is in AdminEventsPanel.',
    '001_shared_schedule.sql': 'Creates the shared active-schedule table and prototype access policies.',
    '002_admin_events.sql': 'Creates individual event rows, imports legacy events, and defines event access and realtime publication setup.',
    '003_student_assistant_schedules.sql': 'Creates shared assistant-profile and solver-result storage.',
    '004_admin_event_assistants.sql': 'Adds assistant identity and display-label fields to event rows.',
    '005_duty_break_setting.sql': 'Adds JSON scheduling settings storage used for duty-gap and workload values.',
    '006_admin_event_descriptions.sql': 'Adds event descriptions without changing the imported class format.',
    '007_schedule_archives.sql': 'Creates the separate archive table, immutable-snapshot guard, and shared prototype policies.',
    '008_csv_scoped_events.sql': 'Adds the schedule key and supporting index that associate each event with a specific uploaded schedule.',
    '009_unique_archive_names.sql': 'Enforces unique normalized names among nondeleted archives. It does not automatically rename conflicting existing snapshots.',
    '010_admin_event_deletions.sql': 'Creates read-only deletion records, a direct-write trigger guard, and an atomic delete RPC. The retained IDs prevent offline event resurrection.',
    'build-backend.mjs': 'Locates Python and runs PyInstaller with the backend import path and OR-Tools collection. Its current onefile output is packaged as an Electron resource.',
    'setup-python.mjs': 'Installs backend requirements during npm installation and reports missing Python or package-installation failures.',
    'register-typescript.mjs': 'Registers a Node loader for testing actual TS and TSX modules without writing a separate test-only implementation.',
    'theme-ui-smoke.cjs': 'Runs isolated hidden Electron UI tests with synthetic local storage and mocked HTTP calls. Optional modes cover assistant edits/removal, booking edits, event deletion, and archive save/rename.',
    'test-backend-process.mjs': 'Tests shutdown commands and failure handling; its optional Windows integration mode launches and terminates only its own disposable process tree.',
    'package.json': 'Lists application dependencies and development, build, test, and platform-packaging commands.',
    'electron-builder.json': 'Defines the app ID, bundled frontend and desktop files, backend resources, copied environment file, icon, and portable/MSI targets.',
    'vite.config.ts': 'Sets the frontend root, relative packaged asset paths, build output, and development server port 5173.',
    'eslint.config.js': 'Configures static code-quality checks for the project.',
    'index.html': 'Provides the root HTML page and root element into which React mounts the interface.',
    'requirements.txt': 'Declares Flask, CORS support, OR-Tools, and PyInstaller dependencies for the Python side.',
}

ROUTE_DESCRIPTIONS = {
    'schedules.py': 'Receives parser requests and shared active-schedule GET and PUT requests; returns validation errors as HTTP 400.',
    'student_assistants.py': 'Receives shared assistant data and solver requests; returns acknowledgement versions needed by local retry queues.',
    'admin_events.py': 'Exposes ordinary event reads, deletion-aware sync snapshots, event PUT, and protected DELETE; distinguishes a deleted event with HTTP 410.',
    'archives.py': 'Exposes archive listing, immutable creation, and deletion; distinguishes a duplicate name with HTTP 409.',
    'health.py': 'Returns an OK status and whether the local backend has usable cloud configuration; it does not prove the cloud project is reachable.',
}
MODEL_DESCRIPTIONS = {
    'schedule.py': 'Validates the shared schedule object and its list fields before storage.',
    'student_assistant.py': 'Checks that shared or solver input is a JSON object. Detailed scheduling validation occurs inside the solver and repositories; this is not a complete nested schema validator.',
    'admin_event.py': 'Checks required event fields, nonblank name and room, schedule key when supplied, and an end time after the start.',
    'archive.py': 'Validates snapshot version, ID consistency, required details, timestamp, schedule lists, assistant data, and viewing dates.',
}
TEST_DESCRIPTIONS = {
    'test_archives.py': 'Checks archive schema, idempotent creation, duplicate names, immutable payload storage, deletion states, and route responses.',
    'test_assistant_rosters.py': 'Checks scoped roster merging and result revisions, including stale inactive caches, offline multi-CSV changes, clears, and older clients.',
    'test_admin_event_deletions.py': 'Checks protected deletes, missing migrations, rejected stale saves, route statuses, pagination, and the migration guard structure.',
    'test_admin_event_descriptions.py': 'Checks description persistence and that missing required schema support is not silently acknowledged.',
    'test_event_schedule.py': 'Checks event schedule keys, assistant persistence, and rejection of invalid scope values.',
    'test_incremental_solver.py': 'Checks preserved old duties, newly added students, remaining slots, and incremental validation rules.',
    'test_schedule_parser.py': 'Checks format detection, shared meetings, metadata, invalid-time TBA behavior, complete official templates, and persisted parsed data.',
    'test_student_assistant_service.py': 'Exercises solver rules and result handling using small deterministic fixtures.',
}

def description(path):
    rel = path.relative_to(ROOT).as_posix()
    if path.name == '__init__.py':
        return 'Marks this folder as a Python package. It contains no separate business process.'
    if '/routes/' in rel:
        return ROUTE_DESCRIPTIONS.get(path.name, '')
    if '/models/' in rel:
        return MODEL_DESCRIPTIONS.get(path.name, '')
    if '/tests/' in rel:
        return TEST_DESCRIPTIONS.get(path.name, 'Regression checks for the feature named by the file.')
    if path.name.startswith('tsconfig'):
        return 'Defines TypeScript compilation scope, strictness, module output, and build relationships for its frontend or Electron target.'
    if path.name in DESCRIPTIONS:
        return DESCRIPTIONS[path.name]
    if path.name.startswith('test-'):
        feature = path.stem.removeprefix('test-').replace('-', ' ')
        caution = ' This separate check may contact a configured service; inspect its fixture and configuration before running it.' if 'database' in feature else ''
        return f'Runs regression checks for {feature}, using the actual related modules and test fixtures.{caution}'
    raise ValueError(f'Missing description for {rel}')

GROUPS = [
    ('Application screens and styling', lambda p: p.as_posix().startswith('frontend/src/ui/')),
    ('API requests storage and formatters', lambda p: p.as_posix().startswith(('frontend/src/api/', 'frontend/src/storage/', 'frontend/src/formatters/', 'frontend/src/export/'))),
    ('Desktop entry points and shared types', lambda p: p.as_posix().startswith('frontend/src/desktop/') or p.as_posix() == 'frontend/src/types.ts'),
    ('Python routes models services and solver', lambda p: p.as_posix().startswith('backend/') and '/tests/' not in p.as_posix()),
    ('Database migrations', lambda p: p.suffix == '.sql'),
    ('Build configuration and verification scripts', lambda p: True),
]

def source_files():
    files = []
    for folder in ['frontend/src', 'backend', 'database/supabase/migrations', 'scripts']:
        for path in (ROOT / folder).rglob('*'):
            if path.suffix in {'.ts', '.tsx', '.css', '.py', '.sql', '.mjs', '.cjs', '.json'} and '__pycache__' not in path.parts:
                files.append(path)
    for rel in ['package.json', 'eslint.config.js', 'frontend/electron-builder.json', 'frontend/vite.config.ts',
                'frontend/index.html', 'frontend/tsconfig.json', 'frontend/tsconfig.app.json', 'frontend/tsconfig.node.json', 'backend/requirements.txt']:
        files.append(ROOT / rel)
    return sorted(set(files), key=lambda p: p.relative_to(ROOT).as_posix())

def symbols(path):
    text = path.read_text(encoding='utf-8-sig')
    found = []
    if path.suffix == '.py':
        tree = ast.parse(text)
        found = [(node.lineno, node.name) for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))]
    elif path.suffix in {'.ts', '.tsx'}:
        for number, line in enumerate(text.splitlines(), 1):
            match = re.match(r'^(?:export\s+(?:default\s+)?)?(?:async\s+)?(?:function|interface|type|class)\s+(\w+)', line)
            if match:
                found.append((number, match.group(1)))
            elif re.match(r'^  const \w+ = (?:async )?\(', line):
                found.append((number, re.search(r'const (\w+)', line).group(1)))
    return found

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.78)
sec.footer_distance = Inches(.25)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Calibri')
    for item in list(style._element.get_or_add_rPr()):
        if item.tag == qn('w:color'):
            item.attrib.pop(qn('w:themeColor'), None)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.size = Pt(27)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next = True
code_style = doc.styles.add_style('Code guide code', 1)
code_style.font.name, code_style.font.size = 'Consolas', Pt(9)
code_style.paragraph_format.space_after = Pt(7)
code_style.paragraph_format.line_spacing = 1.0
code_style.paragraph_format.keep_together = True
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = foot.add_run('Auto Scheduler Code Guide   |   ')
run.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
doc.core_properties.title = 'Auto Scheduler Code Guide'
doc.core_properties.subject = 'Application workflows algorithms and source file reference'
doc.core_properties.author = 'Auto Scheduler Project'

def inline(paragraph, text):
    for index, piece in enumerate(re.split(r'(`[^`]+`)', text)):
        run = paragraph.add_run(piece[1:-1] if piece.startswith('`') and piece.endswith('`') else piece)
        if piece.startswith('`') and piece.endswith('`'):
            run.font.name = 'Consolas'; run.font.size = Pt(9.5)

def paragraph(text, size=None, bold=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = keep
    inline(p, text)
    if size or bold:
        for run in p.runs:
            if size: run.font.size = Pt(size)
            if bold: run.bold = True
    return p

def heading(text, level=1, new_page=False):
    p = doc.add_heading(text, level)
    p.paragraph_format.page_break_before = new_page
    return p

def table(rows):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = {2: [2.4, 4.54], 3: [1.25, 2.35, 3.34],
              4: [0.8, 1.6, 2.5, 2.04], 5: [1.74, 1.3, 1.3, 1.3, 1.3]}[len(rows[0])]
    if rows[0][0] == 'Objective term': widths = [2.2, 1.5, 3.24]
    if rows[0][0] == 'Concept': widths = [1.4, 2.6, 2.94]
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for ri, values in enumerate(rows):
        cells = t.add_row().cells
        for ci, value in enumerate(values):
            cell = cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                element = OxmlElement('w:' + edge); element.set(qn('w:w'), '95'); element.set(qn('w:type'), 'dxa'); margins.append(element)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                element = OxmlElement('w:' + edge); element.set(qn('w:val'), 'single'); element.set(qn('w:sz'), '4'); element.set(qn('w:color'), 'D9D9D9'); borders.append(element)
            tcpr.append(borders)
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'DCE6F1' if ri == 0 else ('F4F6F8' if ri % 2 == 0 else 'FFFFFF')); tcpr.append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            inline(p, value)
            for run in p.runs: run.font.size = Pt(10); run.bold = ri == 0
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

files = source_files()
inventory = []
pending = list(files)
for group, predicate in GROUPS:
    chosen = [path for path in pending if predicate(path.relative_to(ROOT))]
    inventory.append((group, chosen))
    pending = [path for path in pending if path not in chosen]
for path in files:
    description(path)

def emit_inventory():
    paragraph(f'The reference includes {len(files)} source, migration, configuration, and verification files. Empty package markers are included for completeness.')
    for group, paths in inventory:
        heading(group, 2)
        for path in paths:
            rel = path.relative_to(ROOT).as_posix()
            line_count = len(path.read_text(encoding='utf-8-sig').splitlines())
            paragraph(f'{rel}  ({line_count} lines)', size=9.5, bold=True, keep=True)
            paragraph(description(path), size=10.5)

def emit_symbols():
    for path in files:
        rel = path.relative_to(ROOT).as_posix()
        if not rel.startswith(('frontend/src/', 'backend/')) or '/tests/' in rel: continue
        found = symbols(path)
        if not found: continue
        paragraph(rel, size=9.5, bold=True, keep=True).paragraph_format.space_after = Pt(3)
        paragraph('; '.join(f'{name} at line {line}' for line, name in found), size=10).paragraph_format.space_after = Pt(4)

guide = (OUT / 'guide.md').read_text(encoding='utf-8')
start = guide.index('## 6 How the solver builds a schedule')
end = guide.index('## 7 Create schedule', start)
guide = guide[:start] + (OUT / 'cpsat.md').read_text(encoding='utf-8') + '\n\n' + guide[end:]
lines = guide.splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line == '@@inventory': emit_inventory(); i += 1; continue
    if line == '@@symbols': emit_symbols(); i += 1; continue
    if line.startswith('@@code '):
        path, first, last = line[7:].split('|')
        snippet = (ROOT / path).read_text(encoding='utf-8-sig').splitlines()[int(first)-1:int(last)]
        p = doc.add_paragraph('\n'.join(snippet), style='Code guide code')
        i += 1; continue
    if line.startswith('```'):
        block = []; i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            block.append(lines[i]); i += 1
        doc.add_paragraph('\n'.join(block), style='Code guide code'); i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            values = [value.strip() for value in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', value) for value in values): rows.append(values)
            i += 1
        table(rows); continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title'); i += 1; continue
    if line.startswith('## '):
        title = line[3:]
        heading(title, 1, new_page=title != 'Reading route'); i += 1; continue
    if line.startswith('### '): heading(line[4:], 2); i += 1; continue
    block = [line]; i += 1
    while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#', '|', '@@', '```')):
        block.append(lines[i]); i += 1
    text = ' '.join(block)
    paragraph(text, size=9.5 if text.startswith(('Source:', 'Sources:')) else None)

for root in [doc.styles.element, doc.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(DOCX)
manifest = [{ 'file': path.relative_to(ROOT).as_posix(), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
              'lines': len(path.read_text(encoding='utf-8-sig').splitlines()) } for path in files]
(OUT / 'source_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
print(json.dumps({'document': str(DOCX), 'source_files': len(files), 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables)}))
