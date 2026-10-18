from pathlib import Path
import json
from docx import Document
from pypdf import PdfReader
from pdf2image import convert_from_path
from PIL import Image, ImageOps, ImageDraw

base = Path(__file__).resolve().parent
qa = base / 'qa' / 'expanded'
pdf = qa / 'guide.pdf'
document = Document(base / 'Auto Scheduler Code Guide.docx')
reader = PdfReader(pdf)
previous = json.loads((qa / 'page_audit.json').read_text(encoding='utf-8')) if (qa / 'page_audit.json').exists() else []
audit = []
for number, page in enumerate(reader.pages, 1):
    text = page.extract_text() or ''
    audit.append({'page': number, 'characters': len(text),
                  'start': text[:160], 'end': text[-160:]})
(qa / 'page_audit.json').write_text(json.dumps(audit, indent=2), encoding='utf-8')
print('Changed page summaries:', [a['page'] for a in audit if a not in previous])
paths = convert_from_path(str(pdf), dpi=100, output_folder=str(qa), fmt='png',
    paths_only=True, thread_count=4,
    poppler_path=r'C:\Users\jap\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin')
for number, path in enumerate(paths, 1):
    Path(path).replace(qa / f'page-{number}.png')
for start in range(1, len(paths) + 1, 4):
    sheet = Image.new('RGB', (1700, 2260), '#dddddd')
    draw = ImageDraw.Draw(sheet)
    for offset in range(4):
        number = start + offset
        if number > len(paths):
            break
        im = Image.open(qa / f'page-{number}.png').convert('RGB')
        x, y = (offset % 2) * 850, (offset // 2) * 1130
        sheet.paste(im, (x, y + 25))
        draw.text((x + 12, y + 5), f'PAGE {number}', fill='black')
    sheet.save(qa / f'review-{start:02}.png')
print(json.dumps({'pages': len(reader.pages), 'words': sum(len(p.text.split()) for p in document.paragraphs)
    + sum(len(c.text.split()) for t in document.tables for r in t.rows for c in r.cells),
    'tables': len(document.tables), 'short_pages': [a for a in audit if a['characters'] < 350],
    'headings': [p.text for p in document.paragraphs if p.style.name == 'Heading 1']}, indent=2))
